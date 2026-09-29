"""Nạp Danh mục hàng hoá XNK (TT 31/2022) vào `hs_codes`: ghép mô tả các cấp cha vào mô tả dòng 8 số."""

import argparse
import re
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.orm import Session

from app.ai.hs.models import HsCode

# Vị trí cột (0-based) trong sheet đầu tiên của file nguồn; chỉnh theo biên bản kiểm chứng (Task 1.9).
COLUMNS = {"code": 0, "vi": 1, "en": 2}
FIRST_DATA_ROW = 2  # số dòng tiêu đề bỏ qua
SEPARATOR = " > "
EXCLUDED_CHAPTERS = frozenset({98})
_LEADING_DASHES = re.compile(r"^\s*((?:[-–—]\s*)+)")


@dataclass(frozen=True)
class HsRow:
    code: str
    chapter: int
    description_vi: str
    description_en: str | None


@dataclass(frozen=True)
class ImportStats:
    inserted: int
    updated: int
    skipped_ch98: int

    @property
    def total(self) -> int:
        return self.inserted + self.updated


def _normalize_code(raw: Any) -> str:
    return re.sub(r"[.\s]", "", str(raw)) if raw is not None else ""


def _split_level(raw: Any) -> tuple[int, str]:
    """Số dấu gạch đầu dòng chính là cấp; trả (cấp, mô tả đã bỏ gạch)."""
    text_value = re.sub(r"\s+", " ", str(raw)).strip() if raw is not None else ""
    match = _LEADING_DASHES.match(text_value)
    if not match:
        return 0, text_value
    return len(re.findall(r"[-–—]", match.group(1))), text_value[match.end():].strip()


def _join(stack: list[str | None], own: str | None) -> str | None:
    parts = [part for part in [*stack, own] if part]
    return SEPARATOR.join(parts) if own else None


def parse_rows(rows: list[tuple[Any, Any, Any]]) -> tuple[list[HsRow], int]:
    """(code, mô tả VI, mô tả EN) theo thứ tự file → (các dòng mã 8 số của chương 1–97, số dòng chương 98 bị bỏ).

    Nhóm 4 số là cấp 0; dòng phân nhóm (không có mã 8 số) đặt tiêu đề cho cấp bằng số dấu "-" ở đầu mô tả."""
    stack_vi: list[str | None] = []
    stack_en: list[str | None] = []
    out: list[HsRow] = []
    skipped_98 = 0
    for raw_code, raw_vi, raw_en in rows:
        code = _normalize_code(raw_code)
        level, vi = _split_level(raw_vi)
        _, en = _split_level(raw_en)
        if not vi:
            continue
        if len(code) == 4 and code.isdigit():
            level = 0
        elif len(code) not in (0, 6, 8) or (code and not code.isdigit()):
            continue  # tiêu đề chương / phần, hoặc dòng không phải mã
        elif not level and not code:
            continue
        if len(code) == 8:
            if int(code[:2]) in EXCLUDED_CHAPTERS:
                skipped_98 += 1
                continue
            chapter = int(code[:2])
            if not 1 <= chapter <= 97:
                continue
            out.append(HsRow(code, chapter, _join(stack_vi[:level], vi) or vi,
                             _join(stack_en[:level], en) if en else None))
            continue
        del stack_vi[level:]
        del stack_en[level:]
        stack_vi.extend([None] * (level - len(stack_vi)))
        stack_en.extend([None] * (level - len(stack_en)))
        stack_vi.append(vi)
        stack_en.append(en or None)
    return out, skipped_98


def read_sheet(path: Path) -> list[tuple[Any, Any, Any]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.worksheets[0]
        rows = []
        for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            if index <= FIRST_DATA_ROW:
                continue
            width = max(COLUMNS.values()) + 1
            cells = tuple(row) + (None,) * (width - len(row))
            rows.append((cells[COLUMNS["code"]], cells[COLUMNS["vi"]], cells[COLUMNS["en"]]))
        return rows
    finally:
        workbook.close()


def import_hs(db: Session, path: Path) -> ImportStats:
    """Upsert theo `code`; mô tả đổi thì xoá embedding để CLI embed làm lại. Không commit."""
    parsed, skipped_98 = parse_rows(read_sheet(Path(path)))
    existing = {row.code: row for row in db.scalars(select(HsCode))}
    inserted = updated = 0
    for item in parsed:
        current = existing.get(item.code)
        if current is None:
            db.add(HsCode(code=item.code, chapter=item.chapter, description_vi=item.description_vi,
                          description_en=item.description_en))
            inserted += 1
        elif (current.description_vi, current.description_en) != (item.description_vi, item.description_en):
            current.description_vi, current.description_en, current.embedding = (
                item.description_vi, item.description_en, None)
            updated += 1
    db.flush()
    return ImportStats(inserted, updated, skipped_98)


def lookup(db: Session, code: str) -> str:
    row = db.get(HsCode, _normalize_code(code))
    return f"{row.code} | {row.description_vi}" if row else "không có"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Nạp Danh mục TT 31/2022 vào hs_codes")
    parser.add_argument("path", nargs="?", type=Path, help="file .xlsx nguồn")
    parser.add_argument("--lookup", metavar="CODE", help="in mô tả của một mã đã nạp")
    args = parser.parse_args(argv)
    from app.db import SessionLocal  # nạp muộn để test không tạo kết nối

    with SessionLocal() as db:
        if args.lookup:
            print(lookup(db, args.lookup))
            return 0
        if args.path is None:
            parser.error("cần đường dẫn file .xlsx hoặc --lookup <mã>")
        stats = import_hs(db, args.path)
        db.commit()
        total = db.scalar(text("SELECT count(*) FROM hs_codes"))
    print(f"inserted={stats.inserted} updated={stats.updated} skipped_ch98={stats.skipped_ch98} total={total}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
